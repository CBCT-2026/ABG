"""Pre-set format order list for participants 1 to 100 (ABG study).

All 6 orders of the 3 formats (Williams design for 3 treatments), shuffled in blocks of 6,
so orders stay balanced at every multiple of 6 participants. Seed 20261019 (new for ABG).
Order codes match the wisdom tooth list. Run: python3 make_order_list.py
"""
import csv
import random
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
SEED = 20261019
ORDERS = {'A': ('Pamphlet', 'AR', 'Printed'), 'B': ('AR', 'Printed', 'Pamphlet'), 'C': ('Printed', 'Pamphlet', 'AR'),
          'D': ('Pamphlet', 'Printed', 'AR'), 'E': ('Printed', 'AR', 'Pamphlet'), 'F': ('AR', 'Pamphlet', 'Printed')}
HOW = [
    'Give each new participant the next unused participant number, whatever their group or language.',
    'Present the three formats in the order shown on that row. Within each format, show the stages in clinical order '
    '(pre-op, short-term post-op, long-term post-op). The participant enters their number in the form and ticks which '
    'format they saw at each step.',
    f'Random sequence: all 6 orders shuffled in blocks of 6 (seed {SEED}), so the orders stay balanced at every multiple '
    'of 6 participants.',
    'Never skip or reuse a number. If someone stops early, mark Completed "No" and give the next person the next number.',
    'Pilot participants use numbers 901 to 905 with orders A, B, C, D and E, not numbers from this list.',
    'Keep this sheet with the facilitator, not with participants.',
]


def rows():
    rng = random.Random(SEED)
    codes = []
    while len(codes) < 100:
        block = list(ORDERS)
        rng.shuffle(block)
        codes += block
    return [(i, c, *ORDERS[c]) for i, c in enumerate(codes[:100], 1)]


def main():
    data = rows()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Order list'
    head = ['Participant no.', 'Order code', '1st', '2nd', '3rd', 'Language', 'Group', 'Date', 'Completed']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='00674C')
    for r in data:
        ws.append(list(r) + [None] * 4)
    for col, w in zip('ABCDEFGHI', (15, 11, 10, 10, 10, 11, 9, 12, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    how = wb.create_sheet('How to use')
    for line in HOW:
        how.append([line])
    how.column_dimensions['A'].width = 120
    wb.save(HERE / 'participant_order_list.xlsx')
    with open(HERE / 'participant_order_list.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(head[:5])
        w.writerows(data)
    for k in range(6, 101, 6):  # check balance at every multiple of 6
        assert sorted(r[1] for r in data[k - 6:k]) == sorted(ORDERS)
    print('codes 1-12:', ''.join(r[1] for r in data[:12]))


if __name__ == '__main__':
    main()
