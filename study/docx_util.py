"""Small helpers shared by the study Word files (python-docx)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

NOTE = '56636F'


def new_doc(lang='EN', track=False, landscape=False):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Yu Gothic' if lang == 'JA' else 'Calibri'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    if track:  # Review > Track Changes is on when the file opens
        el = OxmlElement('w:trackRevisions')
        doc.settings.element.append(el)
    return doc


def para(doc_or_cell, text='', style=None, bold=False, italic=False, size=None, color=None, justify=True):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hex_fill)
    tc.append(s)


def cell_text(cell, lines, bold_first=False, color=None, size=None):
    """Fill a table cell with one paragraph per line."""
    lines = [lines] if isinstance(lines, str) else lines
    cell.text = ''
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        r = p.add_run(line)
        r.bold = bold_first and i == 0
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if size:
            r.font.size = Pt(size)


def table(doc, header, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.autofit = False
    for g, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        g.set(qn('w:w'), str(int(w * 567)))  # cm to twips
    for c, h in zip(t.rows[0].cells, header):
        cell_text(c, h, bold_first=True)
        shade(c, 'DCE8EF')
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    return t


def add_row(t, values, widths, colors=None):
    cells = t.add_row().cells
    for i, (c, v, w) in enumerate(zip(cells, values, widths)):
        cell_text(c, v, color=(colors or {}).get(i))
        c.width = Cm(w)
    return cells
