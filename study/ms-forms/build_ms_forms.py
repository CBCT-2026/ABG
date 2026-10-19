"""Build the Microsoft Forms files for the ABG study (EN, JA). Adapted from the wisdom tooth study.

Quick import only reliably makes the title and numbered questions (with A. B. C. options).
It merges headings and descriptions into the text, so those are kept out of the import file.
- quick-import_{LANG}.docx: title + numbered questions, for Forms > Quick import.
- form-text_{LANG}.docx: formatted text to paste (form description, sections, Likert rows,
  ranking, thank-you message), each with where it goes. Copy from Word keeps bold and bullets.
Run: python3 build_ms_forms.py   (needs python-docx)
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
LETTERS = 'ABCDEFGHIJKLMNOP'
CAP_EN, CAP_JA = '2 minutes', '2分'  # viewing cap per format (120 s); change both if the pilot cuts it to 90 s

T = {
'EN': dict(
 title='Alveolar Bone Grafting Education Models Study',
 formal='Comparing a 2D pamphlet, an augmented reality 3D model and 3D printed models for explaining '
        'alveolar bone grafting in cleft lip and palate: a within-participant study',
 agree=['I agree', 'I do not agree'],
 langs=['Chinese (Cantonese)', 'Chinese (Mandarin)', 'English', 'French', 'German', 'Hindi', 'Italian', 'Japanese',
        'Korean', 'Portuguese', 'Russian', 'Spanish', 'Vietnamese', 'Other'],
 tools=['2D pamphlet', '3D models on the web page (AR)', '3D printed models'],
 done='All three are done',
 header=[
  ('p', 'A Nagoya University study. It takes about 12 minutes.'),
  ('l', 'Why this study', 'Some children are born with a cleft lip and palate. Many of them also have a gap in the bone '
        'of the upper gum. An operation called alveolar bone grafting fills this gap with new bone, often before the adult '
        'canine tooth comes through. This study compares three ways to explain this operation.'),
  ('l', 'What happens', f'Look at 3 items, one at a time, for up to {CAP_EN} each in a random order: 1) a paper pamphlet, '
        '2) 3D models in augmented reality on a web page (on a study iPad), and 3) 3D printed models. Each item shows 3 stages: '
        'before the operation, soon after it, and years later after braces. After each item, answer a few questions on this '
        'form. At the end, put the 3 items in order. About 12 minutes in total, and no more than 15 minutes.'),
  ('l', 'Free choice', 'Taking part is free choice. Saying no does not change study, work or care at Nagoya University. To stop and '
        'not participate, close this form. Nothing is saved until the "Submit" button is pressed.'),
  ('s', 'Study title: {formal}.'),
 ],
 sections=[
  ('Consent Form', None, [
   ('h', 'Privacy'),
   ('b', 'No name, email, or student or staff number is asked.'),
   ('b', 'Answers have no name on them. So, after "Submit," an answer cannot be found or taken out.'),
   ('b', 'Answers are collected with Microsoft Forms in the Nagoya University Microsoft 365 account, then moved to safe '
         'Nagoya University storage. They are kept for at least 10 years after the results are published, then permanently deleted.'),
   ('b', 'Results are shared only as group results. A data file with no names may be shared for research.'),
   ('h', 'Risks and benefits'),
   ('b', 'The risks are very small. The items show an operation done on children. This may worry some people, such as '
         'parents of a child with a cleft.'),
   ('b', 'This study gives no personal advice. Questions about cleft care are best asked to the cleft care team.'),
   ('b', 'When using AR, stay seated or watch the space around.'),
   ('b', 'There is no direct benefit. The results may help cleft care teams explain this operation more clearly.'),
   ('h', 'Cost'),
   ('b', 'Participation is free.'),
   ('h', 'Study team and contact'),
   ('b', 'Dr Markus Seifert, PhD student, Department of International Medicine Education, School of Medicine, Nagoya University.'),
   ('b', 'Approved by [committee name], Nagoya University (approval number [number]).'),
   ('b', 'For any questions or concerns please email seifert.markus.luke.f5@s.mail.nagoya-u.ac.jp.'),
   ('h', 'Consent'),
   ('b', 'To take part, choose "I agree" for each statement below.'),
  ]),
  ('Study Participant ID', 'c6', []),
  ('Participant Information', 'pnum', [('p', 'A few questions about the participant. No name or identifying information is requested.')]),
  ('2D Pamphlet', 'first', [('p', f'Please look at the paper pamphlet for up to {CAP_EN}. Read the 3 stages in order. '
                            'The staff will say when time is up. Then answer the questions below about the pamphlet.')]),
  ('3D Models on the Web Page (AR)', 'next_pamphlet', [('p', f'Please use the web page on the study iPad for up to {CAP_EN}. '
                            'Look at the 3 stages in order, starting with stage 1. Turn and zoom the models and try "View in AR" '
                            'if wanted. The staff will say when time is up. Then answer the questions below about the web page.')]),
  ('3D Printed Models', 'next_ar', [('p', f'Please look at the 3 printed models and the card for up to {CAP_EN}. Look at the '
                            'stages in order, starting with stage 1. It is fine to pick the models up and take the parts off. '
                            'The staff will say when time is up. Then answer the questions below about the models.')]),
  ('Final Comparison', 'next_print', [('p', 'A few last questions to compare the three items.')]),
 ],
 q=dict(
  pnum='Participant number (the study staff will give this number)',
  consent=['I am 18 years old or older.', 'I have read and understood the study information.',
           'I understand that taking part is free choice, and that I can stop at any time before "Submit."',
           'I understand that my answers have no name on them and cannot be taken back after "Submit."',
           'I agree that my answers, with no name on them, can be published and shared for research.',
           'I agree to take part in this study.'],
  about=[('Which group fits best? (choose the first one that fits)', [
           'Parent or caregiver of a child with a cleft lip and/or palate',
           'Person with a cleft lip and/or palate',
           'Member of the public (no health training)',
           'Dental or medical student',
           'Dentist or orthodontist',
           'Oral and maxillofacial surgeon, cleft or plastic surgeon, or other medical doctor',
           'Other healthcare worker (for example nurse, speech pathologist, dental hygienist, technician or other)']),
         ('Age group', ['18-24', '25-34', '35-44', '45-54', '55-64', '65 or older']),
         ('Gender', ['Man', 'Woman', 'Non-binary', 'Prefer not to say', 'Other']),
         ('First language (the language learned first as a child)', 'LANGS'),
         ('Language(s) spoken the most in daily life', 'LANGS'),
         ('Any personal link to cleft lip and/or palate? (choose all that fit)', [
           'Born with a cleft', 'Parent or caregiver of a child with a cleft', 'Other family member or friend with a cleft',
           'Work with people who have a cleft', 'None', 'Prefer not to say']),
         ('Ever used augmented reality (AR) on a phone or tablet before?', ['Yes', 'No', 'Not sure'])],
  first='Which item is first? (the study staff will say)',
  nxt='Which item is next? (the study staff will say)',
  ar=('Did the models open in AR (shown in the room)?', ['Yes, it worked', 'I tried, but it did not work', 'I did not try to open it']),
  pref='If this operation were planned for me or my child, which item would be best to explain it?',
  f3='Any comments? (optional)',
 ),
 ueq_title='For each pair of words, choose the circle that best matches the first impression of this item. '
           'There are no right or wrong answers.',
 ueq=['obstructive ... supportive', 'complicated ... easy', 'inefficient ... efficient', 'confusing ... clear',
      'boring ... exciting', 'not interesting ... interesting', 'conventional ... inventive', 'usual ... leading edge'],
 und_title='Choose one answer for each statement about this item.',
 und=['This helped me understand where the gap in the bone is and what the bone graft fills.',
      'This helped me understand why the bone graft is done before the adult canine tooth comes through.',
      'This helped me understand what happens after the operation, soon after and later with braces.'],
 und_cols='Strongly disagree, Disagree, Slightly disagree, Neutral, Slightly agree, Agree, Strongly agree',
 rank='Please put the three items in order, from most helpful (top) to least helpful (bottom), for understanding '
      'alveolar bone grafting.',
 thanks='Many thanks for taking part. For questions about cleft care, please talk to a cleft care team.',
),
'JA': dict(
 title='歯そう骨移植の教育モデルの研究',
 formal='口唇口蓋裂の歯槽骨移植の説明における、2Dパンフレット、拡張現実（AR）3Dモデル、3Dプリントモデルの比較：参加者内比較研究',
 agree=['同意する', '同意しない'],
 langs=['中国語（広東語）', '中国語（標準語）', '英語', 'フランス語', 'ドイツ語', 'ヒンディー語', 'イタリア語', '日本語',
        '韓国語', 'ポルトガル語', 'ロシア語', 'スペイン語', 'ベトナム語', 'その他'],
 tools=['2Dパンフレット', 'ウェブページ上の3Dモデル（AR）', '3Dプリントモデル'],
 done='3つとも終わった',
 header=[
  ('p', '名古屋大学の研究です。約12分かかります。'),
  ('l', 'この研究について', '生まれつき、口しん口がいれつ（くちびるや上あごのわれ目）がある子どもがいます。その多くは、上の歯ぐきの'
        '骨にもすき間があります。歯そう骨移植は、このすき間を新しい骨でうめる手術です。大人の犬歯が生えてくる前に行うことが'
        '多いです。この研究では、この手術を説明する3つの方法をくらべます。'),
  ('l', '参加の内容', f'3つの資料を、ランダムな順番で1つずつ、それぞれ{CAP_JA}まで見ます：1）紙のパンフレット、2）ウェブページ上の'
        'AR（拡張現実）の3Dモデル（研究用のiPadで見ます）、3）3Dプリントモデル。どの資料も、手術の前、手術のすぐ後、'
        'きょう正治りょうの後の数年後、の3つの時期を表しています。1つ見るごとに、このフォームでいくつかの質問に答えます。'
        '最後に、3つに順位を付けます。全部で約12分、長くても15分以内です。'),
  ('l', '参加は自由です', '参加しなくても、名古屋大学での学業、仕事、しんりょうに、えいきょうはありません。参加をやめるときは、'
        'このフォームを閉じてください。「送信」ボタンをおすまで、回答は保存されません。'),
  ('s', '研究名：{formal}'),
 ],
 sections=[
  ('同意書', None, [
   ('h', '回答のあつかい'),
   ('b', '名前、メールアドレス、学生番号、職員番号は聞きません。'),
   ('b', '回答には名前が付いていません。そのため「送信」のあとは、回答を見つけたり取り消したりできません。'),
   ('b', '回答は名古屋大学のMicrosoft 365アカウントのMicrosoft Formsで集め、名古屋大学の安全な保存場所に移します。'
         '結果の発表から少なくとも10年間保存し、そのあと完全に消去します。'),
   ('b', '結果は全体の結果としてのみ発表します。名前のない回答データを、研究のために共有することがあります。'),
   ('h', '危険と利益'),
   ('b', '危険はほとんどありません。資料では、子どもが受ける手術を説明しています。口しん口がいれつのある子どもの家族など、'
         '不安に感じる人もいるかもしれません。'),
   ('b', 'この研究では、一人一人への助言はしません。口しん口がいれつの治りょうについての質問は、担当の医りょうチームに相談してください。'),
   ('b', 'ARを使うときは、座ったまま、または周りに気を付けてください。'),
   ('b', '参加した人への直接の利益はありません。結果は、医りょうチームがこの手術をより分かりやすく説明するのに役立つかもしれません。'),
   ('h', '費用'),
   ('b', '参加は無料です。'),
   ('h', '研究をする人・問い合わせ先'),
   ('b', 'Markus Seifert（医師、名古屋大学 医学部 International Medicine Education 博士課程学生）'),
   ('b', '名古屋大学［委員会名］の承認を受けています（承認番号：［番号］）。'),
   ('b', '質問や意見は、メールでお問い合わせください：seifert.markus.luke.f5@s.mail.nagoya-u.ac.jp'),
   ('h', '同意'),
   ('b', '参加する場合は、下のすべての文で「同意する」を選んでください。'),
  ]),
  ('参加者番号', 'c6', []),
  ('参加者について', 'pnum', [('p', '参加者についての質問です。名前や、個人が分かる情報は聞きません。')]),
  ('2Dパンフレット', 'first', [('p', f'紙のパンフレットを{CAP_JA}まで見てください。3つの時期を順番に読んでください。'
                              '時間になったら担当者が知らせます。そのあと、パンフレットについての質問に答えてください。')]),
  ('ウェブページ上の3Dモデル（AR）', 'next_pamphlet', [('p', f'研究用のiPadで、ウェブページを{CAP_JA}まで使ってください。'
                              '3つの時期を、1から順番に見てください。モデルを回したり大きくしたりして、よければ「AR表示」も'
                              '試してください。時間になったら担当者が知らせます。そのあと、ウェブページについての質問に答えてください。')]),
  ('3Dプリントモデル', 'next_ar', [('p', f'3つのプリントモデルとカードを{CAP_JA}まで見てください。時期の1から順番に見てください。'
                              'モデルを手に取ったり、部品を外したりしてかまいません。時間になったら担当者が知らせます。'
                              'そのあと、モデルについての質問に答えてください。')]),
  ('最後に：3つをくらべる', 'next_print', [('p', '3つの資料をくらべる、最後の質問です。')]),
 ],
 q=dict(
  pnum='参加者番号（担当者が伝えます）',
  consent=['18さい以上です。', '研究の説明を読み、理解しました。',
           '参加は自由であり、「送信」の前ならいつでもやめられることを理解しました。',
           '回答には名前が付かず、「送信」のあとは取り消せないことを理解しました。',
           '名前のない回答が、研究のために発表・共有されることに同意します。',
           'この研究に参加することに同意します。'],
  about=[('一番当てはまるグループはどれですか？（当てはまる最初のものを選んでください）', [
           '口しん口がいれつのある子どもの保護者',
           '口しん口がいれつのある本人',
           '医りょう系の資格・教育のない方',
           '歯学部生・医学部生',
           '歯科医師・きょう正歯科医',
           '口くう外科医、形成外科医、その他の医師',
           'その他の医りょうの仕事（看護師、言語ちょうかく士、歯科衛生士、技師など）']),
         ('年れい層', ['18～24さい', '25～34さい', '35～44さい', '45～54さい', '55～64さい', '65さい以上']),
         ('性別', ['男性', '女性', 'ノンバイナリー', '答えない', 'その他']),
         ('最初に身に付けた言語（子どものころ最初に覚えた言葉）', 'LANGS'),
         ('ふだんの生活で一番よく話す言語（いくつ選んでもかまいません）', 'LANGS'),
         ('口しん口がいれつと関わりがありますか？（当てはまるものをすべて選んでください）', [
           '生まれつき口しん口がいれつがある', '口しん口がいれつのある子どもの保護者',
           '家族や友人に口しん口がいれつのある人がいる', '仕事で口しん口がいれつのある人に関わっている', 'ない', '答えない']),
         ('スマートフォンやタブレットで拡張現実（AR）を使ったことはありますか？', ['はい', 'いいえ', '分からない'])],
  first='最初の資料はどれですか？（担当者が伝えます）',
  nxt='次の資料はどれですか？（担当者が伝えます）',
  ar=('モデルはAR表示（まわりの空間に表示）で開きましたか？', ['はい、表示できた', '試したが、うまくいかなかった', '開こうとしなかった']),
  pref='もし自分や自分の子どもがこの手術を受けるとしたら、どの資料で説明してほしいですか？',
  f3='ご意見があれば書いてください（任意）',
 ),
 ueq_title='それぞれの言葉の組について、この資料の印象に一番近いところを選んでください。正解・不正解はありません。直感で答えてください。',
 ueq=['妨げになる ... 助けられる', '複雑 ... 簡単', '効率が悪い ... 効率が良い', 'ごちゃごちゃしている ... すっきりしている',
      '退屈だ ... エキサイティングだ', 'おもしろくない ... おもしろい', '従来どおり ... 独特だ', '普通 ... 斬新的'],
 und_title='この資料についての各文に、どのくらいそう思うかを選んでください。',
 und=['この資料で、骨のすき間がどこにあり、骨移植で何をうめるのかが理解できた。',
      'この資料で、大人の犬歯が生えてくる前に骨移植をする理由が理解できた。',
      'この資料で、手術のすぐ後と、その後のきょう正治りょうまでに、何が起こるかが理解できた。'],
 und_cols='全くそう思わない、そう思わない、あまりそう思わない、どちらともいえない、ややそう思う、そう思う、とてもそう思う',
 rank='歯そう骨移植を理解するのに、一番役に立つもの（上）から一番役に立たないもの（下）の順に、3つを並べてください。',
 thanks='ご参加ありがとうございました。口しん口がいれつの治りょうについて質問がある場合は、担当の医りょうチームに相談してください。',
),
}


def questions(t):
    """Imported questions in form order: (key, text, options or None). Likerts and the ranking are added by hand."""
    q = t['q']
    tools = t['tools']
    out = [(f'c{i+1}', s, t['agree']) for i, s in enumerate(q['consent'])]
    out.append(('pnum', q['pnum'], None))
    for i, (s, o) in enumerate(q['about']):
        out.append((f'p{i+1}', s, t['langs'] if o == 'LANGS' else o))
    out.append(('first', q['first'], tools))
    out.append(('next_pamphlet', q['nxt'], [tools[1], tools[2], t['done']]))
    out.append(('ar', q['ar'][0], q['ar'][1]))
    out.append(('next_ar', q['nxt'], [tools[0], tools[2], t['done']]))
    out.append(('next_print', q['nxt'], [tools[0], tools[1], t['done']]))
    out.append(('pref', q['pref'], tools))
    out.append(('f3', q['f3'], None))
    return out


def new_doc(lang):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Yu Gothic' if lang == 'JA' else 'Calibri'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    st.font.size = Pt(11)
    return doc


def para(doc, text, style=None, bold=False, italic=False, size=None, color=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def rich(doc, blocks, t):
    for kind, s, *extra in blocks:
        s = s.format(formal=t['formal'])
        if kind == 'l':
            p = para(doc, '')
            p.add_run(s + ': ').bold = True
            p.add_run(extra[0])
        elif kind == 'h':
            para(doc, s, bold=True)
        elif kind == 'b':
            para(doc, s, style='List Bullet')
        elif kind == 's':
            para(doc, s, italic=True, size=9)
        else:
            para(doc, s)


def build_import(lang):
    t = T[lang]
    doc = new_doc(lang)
    para(doc, t['title'], style='Title')
    for n, (_, s, opts) in enumerate(questions(t), 1):
        para(doc, f'{n}. {s}')
        for i, o in enumerate(opts or []):
            para(doc, f'{LETTERS[i]}. {o}')
        para(doc, '')
    doc.save(HERE / f'quick-import_{lang}.docx')


NOTE = '0F6E56'


def build_text(lang):
    t = T[lang]
    qs = questions(t)
    num = {k: i for i, (k, _, _) in enumerate(qs, 1)}
    doc = new_doc(lang)
    para(doc, f'Text to paste into Microsoft Forms ({lang})', style='Title')
    para(doc, 'Green lines say where each piece goes; copy only the black text. Copying from Word keeps bold and bullets.',
         color=NOTE)
    para(doc, '1. Form description (click the title area, then the description box; replace what is there)',
         style='Heading 1')
    rich(doc, t['header'], t)
    for i, (title, k, blocks) in enumerate(t['sections']):
        where = 'at the start, above question 1' if k is None else f'below question {num[k]} of the import'
        para(doc, f'{i + 2}. Section: {title} ({where})', style='Heading 1')
        para(doc, 'Section title:', color=NOTE)
        para(doc, title, bold=True)
        if blocks:
            para(doc, 'Section description:', color=NOTE)
            rich(doc, blocks, t)
    para(doc, '9. Two Likert questions for each of the three item sections (the AR question goes after them in the web page section)', style='Heading 1')
    para(doc, 'Likert A (UEQ-S). Question title (this text is what participants see):', color=NOTE)
    para(doc, t['ueq_title'])
    para(doc, 'Statements (one per row). Official UEQ-S items: do not change the wording or order:', color=NOTE)
    for s in t['ueq']:
        para(doc, s)
    para(doc, 'Options (columns): 1, 2, 3, 4, 5, 6, 7', color=NOTE)
    para(doc, 'Likert B (understanding). Question title (this text is what participants see):', color=NOTE)
    para(doc, t['und_title'])
    para(doc, 'Statements:', color=NOTE)
    for s in t['und']:
        para(doc, s)
    para(doc, 'Options (columns): ' + t['und_cols'], color=NOTE)
    para(doc, f'10. Ranking question, top of the final section (above question {num["pref"]})', style='Heading 1')
    para(doc, t['rank'])
    para(doc, 'Options: ' + ' / '.join(t['tools']), color=NOTE)
    para(doc, '11. Thank-you message (Settings > Customize thank you message)', style='Heading 1')
    para(doc, t['thanks'])
    doc.save(HERE / f'form-text_{lang}.docx')


def all_text(lang):
    """Every participant-facing string, for the wording check."""
    t = T[lang]
    out = [t['title'], t['done'], t['ueq_title'], t['und_title'], t['und_cols'], t['rank'], t['thanks'], *t['agree'], *t['tools'], *t['und']]
    for kind, s, *extra in t['header']:
        if kind != 's':
            out += [s, *extra]
    for title, _, blocks in t['sections']:
        out.append(title)
        out += [b[1] for b in blocks]
    for _, s, opts in questions(t):
        out += [s, *(o for o in opts or [] if o not in t['langs'])]
    return out


if __name__ == '__main__':
    for lang in T:
        build_import(lang)
        build_text(lang)
        print(lang, [(i, k) for i, (k, _, _) in enumerate(questions(T[lang]), 1)])
